"""Презентация для участников марафона: PPTX из реальных экранов бота.

Скриншоты берём из data/shots (их рисует scripts/make_shots.py).
Запуск:  python -m scripts.make_shots && python -m scripts.make_deck
Результат: data/Марафон_добрых_дел_для_участников.pptx
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "data" / "shots"
OUT = ROOT / "data" / "Марафон_добрых_дел_для_участников.pptx"

BG = RGBColor(0x17, 0x21, 0x2B)        # фон Telegram dark
CARD = RGBColor(0x1E, 0x2C, 0x3A)
TEXT = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x94, 0xA7, 0xB8)
ACCENT = RGBColor(0x5E, 0xA9, 0xEE)
GOLD = RGBColor(0xF2, 0xB8, 0x3E)

W, H = Inches(13.333), Inches(7.5)
FONT = "Segoe UI"


def _bg(slide) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def _box(slide, x, y, w, h, *, fill=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.adjustments[0] = 0.06
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill or CARD
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _text(slide, x, y, w, h, blocks, *, align=PP_ALIGN.LEFT):
    """blocks: список (текст, размер, цвет, жирный, отступ_сверху_pt)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (txt, size, color, bold, space) in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(space)
        p.line_spacing = 1.15
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
    return tb


def _shots(slide, names, x, y, w, h, gap=Inches(0.25)):
    """Вписать один-два скриншота в прямоугольник (x, y, w, h), не выходя за его границы."""
    from PIL import Image
    ratios = []
    for name in names:
        with Image.open(SHOTS / f"{name}.png") as im:
            ratios.append(im.width / im.height)
    free = w - gap * (len(names) - 1)
    # Высота, при которой суммарная ширина влезает в free, но не больше h.
    height = min(h, int(free / sum(ratios)))
    total = int(height * sum(ratios)) + gap * (len(names) - 1)
    cx = x + (w - total) // 2
    cy = y + (h - height) // 2
    for name, ratio in zip(names, ratios):
        pw = Emu(int(height * ratio))
        slide.shapes.add_picture(str(SHOTS / f"{name}.png"), cx, cy, width=pw, height=Emu(int(height)))
        cx = cx + pw + gap


def _title_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _text(s, Inches(1.1), Inches(2.0), Inches(7.6), Inches(3.5), [
        ("МАРАФОН", 20, ACCENT, True, 0),
        ("Добрых дел", 60, TEXT, True, 6),
        ("3 недели · 12 добрых дел · командный зачёт", 22, MUTED, False, 18),
        ("Всё происходит в Telegram-боте: получаешь задание, делаешь доброе дело, "
         "присылаешь фото — получаешь баллы себе и своей команде.", 18, TEXT, False, 22),
    ])
    _shots(s, ["06_menu"], Inches(8.7), Inches(0.6), Inches(4.0), Inches(6.3))


def _agenda(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _text(s, Inches(0.9), Inches(0.7), Inches(11.5), Inches(1.0), [
        ("Что вас ждёт", 40, TEXT, True, 0),
    ])
    items = [
        ("1", "3 недели марафона", "Каждую неделю открывается 4 новых задания."),
        ("2", "12 добрых дел", "От благодарности коллеге до помощи приюту."),
        ("3", "200–500 баллов", "Чем сложнее дело, тем дороже оно стоит."),
        ("4", "Командный зачёт", "Ваши баллы идут вам и вашей команде."),
        ("5", "Рейтинг", "Видно свой вклад и место команды в реальном времени."),
        ("6", "Проверка P&C", "Каждый отчёт смотрит человек, а не робот."),
    ]
    x0, y0 = Inches(0.9), Inches(1.9)
    cw, ch = Inches(3.6), Inches(1.55)
    for i, (num, head, sub) in enumerate(items):
        cx = x0 + (cw + Inches(0.35)) * (i % 3)
        cy = y0 + (ch + Inches(0.35)) * (i // 3)
        _box(s, cx, cy, cw, ch)
        _text(s, cx + Inches(0.3), cy + Inches(0.22), cw - Inches(0.6), ch, [
            (num, 13, GOLD, True, 0),
            (head, 19, TEXT, True, 2),
            (sub, 13, MUTED, False, 4),
        ])


def _weeks(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _text(s, Inches(0.9), Inches(0.7), Inches(11.5), Inches(1.0), [
        ("Три недели — три темы", 40, TEXT, True, 0),
        ("Неделя открывается организаторами. Пока неделя не открыта, задания скрыты.",
         16, MUTED, False, 8),
    ])
    weeks = [
        ("Неделя 1", ["Выразить благодарность коллеге — 200",
                      "Тайный доброжелатель — 400",
                      "Поделиться обучающим материалом — 300",
                      "Визит в детский дом / хоспис — 500"]),
        ("Неделя 2", ["Вклад в корпоративную библиотеку — 200",
                      "Экологический сбор — 400",
                      "Настольные игры с коллегами — 300",
                      "World CleanUp Day — 500"]),
        ("Неделя 3", ["Покормить животных в приюте — 500",
                      "Сбор одежды на благотворительность — 400",
                      "Вовлечённость и благополучие коллег — 500",
                      "Сходить в кино с коллегами — 300"]),
    ]
    x0, y0 = Inches(0.9), Inches(2.15)
    cw, ch = Inches(3.75), Inches(4.4)
    for i, (head, rows) in enumerate(weeks):
        cx = x0 + (cw + Inches(0.4)) * i
        _box(s, cx, y0, cw, ch)
        total = sum(int(r.rsplit("—", 1)[1]) for r in rows)
        from app.config import settings as cfg

        period = cfg.weeks[i].label.replace("–", "—")
        blocks = [(head, 24, GOLD, True, 0),
                  (f"{period} · {total} баллов", 13, MUTED, False, 2)]
        for r in rows:
            blocks.append(("•  " + r, 15, TEXT, False, 14))
        _text(s, cx + Inches(0.32), y0 + Inches(0.35), cw - Inches(0.64), ch, blocks)


def _steps_slide(prs, step, title, subtitle, bullets, shots, *, note=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _text(s, Inches(0.9), Inches(0.62), Inches(6.4), Inches(1.6), [
        (step, 15, ACCENT, True, 0),
        (title, 36, TEXT, True, 4),
        (subtitle, 16, MUTED, False, 8),
    ])
    blocks = []
    for i, b in enumerate(bullets):
        blocks.append((b, 17, TEXT, False, 0 if i == 0 else 16))
    _text(s, Inches(0.9), Inches(2.55), Inches(6.2), Inches(3.6), blocks)
    if note:
        _box(s, Inches(0.9), Inches(6.0), Inches(6.2), Inches(0.95))
        _text(s, Inches(1.15), Inches(6.2), Inches(5.7), Inches(0.7),
              [(note, 14, MUTED, False, 0)])
    _shots(s, shots, Inches(7.45), Inches(0.55), Inches(5.5), Inches(6.4))


def _closing(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _text(s, Inches(0.9), Inches(1.4), Inches(7.2), Inches(4.5), [
        ("Коротко: что делать", 40, TEXT, True, 0),
        ("1.  Открыть бота и нажать «Начать»", 20, TEXT, False, 22),
        ("2.  Пройти регистрацию: имя → номер → команда", 20, TEXT, False, 12),
        ("3.  Дождаться подтверждения от P&C", 20, TEXT, False, 12),
        ("4.  Открыть «Задания недели» и выбрать дело", 20, TEXT, False, 12),
        ("5.  Сделать доброе дело и прислать отчёт по шагам", 20, TEXT, False, 12),
        ("6.  Получить баллы и следить за рейтингом", 20, TEXT, False, 12),
        ("Вопросы — раздел «Помощь» в боте.", 17, MUTED, False, 26),
    ])
    _shots(s, ["14_help"], Inches(8.6), Inches(0.7), Inches(4.1), Inches(6.1))


def build() -> Path:
    if not SHOTS.exists() or not any(SHOTS.glob("*.png")):
        subprocess.run([sys.executable, "-m", "scripts.make_shots"], cwd=ROOT, check=True)

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    _title_slide(prs)
    _agenda(prs)
    _weeks(prs)

    _steps_slide(
        prs, "ШАГ 1 · РЕГИСТРАЦИЯ", "Знакомимся",
        "Займёт минуту, делается один раз.",
        ["Откройте бота и нажмите «Начать».",
         "Напишите имя и фамилию одним сообщением.",
         "Отправьте номер телефона: кнопкой «Поделиться номером» "
         "или текстом в формате +998 90 123 45 67.",
         "Выберите свою команду из списка.",
         "Пришли по ссылке-приглашению от коллеги? Команда уже выбрана — шагов всего два."],
        ["01_welcome", "02_reg_name"],
        note="Кнопка «Поделиться номером» работает в мобильном Telegram. "
             "В десктопной версии просто напишите номер текстом.",
    )

    _steps_slide(
        prs, "ШАГ 2 · ПОДТВЕРЖДЕНИЕ", "Ждём зелёный свет",
        "Заявку проверяет сотрудник P&C.",
        ["После регистрации заявка уходит на проверку.",
         "Бот пришлёт сообщение, когда вас подтвердят.",
         "До подтверждения задания недоступны — это нормально.",
         "Если ждёте дольше рабочего дня — напишите в «Помощь»."],
        ["03_reg_phone", "05_pending"],
    )

    _steps_slide(
        prs, "ШАГ 3 · ГЛАВНОЕ МЕНЮ", "Ваш штаб",
        "Отсюда доступно всё.",
        ["«Задания недели» — список дел текущей недели.",
         "«Моя команда» — состав и баллы команды.",
         "«Мой вклад» — что вы сдали и что зачтено.",
         "«Рейтинг» — таблица команд.",
         "«Правила» и «Помощь» — если что-то непонятно."],
        ["06_menu", "13_rating"],
        note="«Моя команда» → «Пригласить в команду» — ссылка, по которой коллега "
             "попадёт именно в вашу команду и зарегистрируется в два шага.",
    )

    _steps_slide(
        prs, "ШАГ 4 · ВЫБОР ЗАДАНИЯ", "Выбираем доброе дело",
        "Задания открываются по неделям.",
        ["Нажмите «Задания недели» — увидите 4 дела и их стоимость.",
         "Откройте любое задание кнопкой «Задание 1…4».",
         "Прочитайте описание: что нужно сделать и сколько это стоит.",
         "Нажмите «Выполнил» — бот проведёт по шагам."],
        ["07_week", "08_task"],
        note="Если недели ещё не открыты, бот честно скажет: задания скоро будут доступны.",
    )

    _steps_slide(
        prs, "ШАГ 5 · ОТЧЁТ ПО ШАГАМ", "Присылаем доказательства",
        "Каждый шаг — отдельное сообщение.",
        ["Бот присылает шаг: что сфотографировать или написать.",
         "В каждом шаге есть блок «Что должно быть видно» и «Что не подойдёт».",
         "Отправьте фото или текст ответным сообщением.",
         "Бот подтвердит шаг и откроет следующий.",
         "Когда шаги закончились — нажмите «Отправить на проверку»."],
        ["09_step1", "10_step2"],
        note="Отчёт можно продолжить позже — бот помнит, на каком шаге вы остановились.",
    )

    _steps_slide(
        prs, "ШАГ 6 · БАЛЛЫ", "Получаем результат",
        "Отчёт смотрит P&C.",
        ["После отправки отчёт уходит на проверку.",
         "Если всё в порядке — баллы падают вам и команде, придёт уведомление.",
         "Если чего-то не хватило — бот объяснит, что поправить, и можно переслать.",
         "Свои баллы всегда видно в «Мой вклад», командные — в «Рейтинге»."],
        ["11_ready", "13_rating"],
    )

    _closing(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(f"Готово: {build()}")
