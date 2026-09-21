"""Excel export for P&C (participants, teams, submissions, points log)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select

from . import services
from .models import PointsLog, User


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, name="Arial")
    for r in rows:
        ws.append(r)
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)


async def export_survey_xlsx(s, path: Path, survey: str | None = None) -> Path:
    """Ответы опроса отдельным файлом: в Telegram такой список читать невозможно."""
    survey = survey or services.SURVEY_CODE
    report = await services.survey_report(s, survey)
    g_labels = dict(services.SURVEY_QUESTIONS["gender"]["options"])
    p_labels = dict(services.SURVEY_QUESTIONS["psy"]["options"])

    wb = Workbook()
    wb.remove(wb.active)

    rows = []
    for item in report["finished"] + report["partial"]:
        u = item["user"]
        answers = item["answers"]
        rows.append([
            u.full_name or u.display_name,
            f"@{u.username}" if u.username else "",
            u.department or "",
            f"{u.team.emoji} {u.team.name}" if u.team else "",
            g_labels.get(answers.get("gender", ""), ""),
            p_labels.get(answers.get("psy", ""), ""),
            "полностью" if services.survey_next_question(answers) is None else "не закончил",
            item["at"].strftime("%d.%m.%Y %H:%M") if item.get("at") else "",
            u.tg_id,
        ])
    rows.sort(key=lambda r: str(r[0]).lower())
    _sheet(wb, "Ответы",
           ["ФИО", "Username", "Отдел", "Команда", "Пол", "Консультация психолога",
            "Статус", "Когда ответил", "TG id"], rows)

    finished = len(report["finished"])
    summary = [
        ["Участников в марафоне", report["invited"]],
        ["Прошли опрос полностью", finished],
        ["Начали и не закончили", len(report["partial"])],
        ["Не открывали", max(report["invited"] - finished - len(report["partial"]), 0)],
        ["", ""],
    ]
    for code, label in services.SURVEY_QUESTIONS["gender"]["options"]:
        n = report["by_gender"].get(code, 0)
        summary.append([label, n])
    summary.append(["", ""])
    for code, label in services.SURVEY_QUESTIONS["psy"]["options"]:
        n = report["by_psy"].get(code, 0)
        summary.append([f"Психолог: {label}", n])
    _sheet(wb, "Сводка", ["Показатель", "Значение"], summary)

    cross = []
    for g_code, g_label in services.SURVEY_QUESTIONS["gender"]["options"]:
        total_g = report["by_gender"].get(g_code, 0)
        yes = report["cross"].get((g_code, "yes"), 0)
        cross.append([g_label, total_g, yes, total_g - yes,
                      f"{round(yes * 100 / total_g)}%" if total_g else "—"])
    _sheet(wb, "Пол × ответ",
           ["Пол", "Всего ответили", "Актуально", "Не актуально", "Доля «актуально»"], cross)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


async def export_xlsx(s, path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    rows = await services.leaderboard(s)
    _sheet(
        wb,
        "Команды",
        ["Место", "Команда", "Баллы", "Участников", "Состав"],
        [[i, f"{r['team'].emoji} {r['team'].name}", r["points"], len(r["members"]), ", ".join(m.display_name for m in r["members"])] for i, r in enumerate(rows, 1)],
    )

    users = await services.list_participants(s)
    upts = await services.user_points_map(s)
    _sheet(
        wb,
        "Участники",
        ["ФИО", "Username", "Телефон", "TG id", "Отдел", "Город", "Команда", "Статус", "Баллы", "Модерация (кто)", "Модерация (когда)", "Причина отклонения", "Причина дисквалификации"],
        [
            [
                u.full_name,
                u.username,
                u.phone,
                u.tg_id,
                u.department,
                u.city,
                u.team.name if u.team else "",
                u.status.value,
                upts.get(u.id, 0),
                u.moderated_by,
                u.moderated_at.strftime("%d.%m.%Y %H:%M") if u.moderated_at else "",
                u.reject_reason,
                u.disqualified_reason,
            ]
            for u in users
        ],
    )

    subs = []
    for u in users:
        subs.extend(await services.user_submissions(s, u.id))
    _sheet(
        wb,
        "Отчёты",
        ["ID", "Участник", "Команда", "Неделя", "№ задания", "Задание", "Опция", "Статус", "Баллы", "Файлов", "Заметка", "Комментарий P&C", "Отправлен", "Проверен"],
        [
            [
                x.id,
                x.user.display_name,
                x.user.team.name if x.user.team else "",
                x.week,
                x.task.code,
                x.task.title,
                x.option.title if x.option else "",
                x.status.value,
                x.points_awarded,
                len(x.files or []),
                x.note,
                x.review_comment,
                x.submitted_at.strftime("%d.%m.%Y %H:%M") if x.submitted_at else "",
                x.reviewed_at.strftime("%d.%m.%Y %H:%M") if x.reviewed_at else "",
            ]
            for x in subs
        ],
    )

    logs = (await s.execute(select(PointsLog, User).join(User, User.id == PointsLog.user_id).order_by(PointsLog.created_at))).all()
    _sheet(
        wb,
        "Журнал баллов",
        ["Дата", "Участник", "Изменение", "Причина", "Кто (tg id)"],
        [[l.created_at.strftime("%d.%m.%Y %H:%M"), u.display_name, l.delta, l.reason, l.actor_tg_id] for l, u in logs],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
